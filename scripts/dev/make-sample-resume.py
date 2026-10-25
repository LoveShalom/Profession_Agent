import os

from docx import Document

os.makedirs("data/tmp", exist_ok=True)

doc = Document()
doc.add_paragraph("张三 · 后端开发工程师")
doc.add_paragraph("手机：138-0000-0000  邮箱：zhangsan@example.com")

doc.add_paragraph("工作经历")
doc.add_paragraph("某某科技有限公司  后端开发实习生")
doc.add_paragraph("2024.06 - 2024.09")
doc.add_paragraph("参与订单系统开发。背景：大促期间系统频繁超时。")
doc.add_paragraph("任务：把 P99 延迟降到 200ms 以内。")
doc.add_paragraph("行动：重写慢查询 SQL，引入 Redis 缓存，按用户分库分表。")
doc.add_paragraph("结果：P99 从 800ms 降到 180ms，吞吐提升 30%。")

doc.add_paragraph("某某大学数据实验室  科研助理")
doc.add_paragraph("2023.09 - 2024.05")
doc.add_paragraph("负责日志清洗脚本，使用 Python 处理 TB 级日志。")
doc.add_paragraph("结果：处理速度提升 40%。")

doc.add_paragraph("页脚占位")
doc.add_paragraph("页脚占位")
doc.add_paragraph("页脚占位")

out = "data/tmp/resume-sample.docx"
doc.save(out)
print(f"已生成 DOCX：{out}（{os.path.getsize(out)} 字节）")
